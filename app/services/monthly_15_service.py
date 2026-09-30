# -*- coding: utf-8 -*-
import os
import sys
import json
import sqlite3
import shutil
import re
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(os.path.dirname(CURRENT_DIR))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from app.db import get_db, log_audit

MASTER_EXCEL_PATH = r'e:\мониторинг\Фарғона_вилояти_15_талик_2026_йил_8_ойи_2_1.xlsx'
EXPORTS_DIR = os.path.join(PROJECT_ROOT, "exports")
BASE_STORAGE_DIR = r"E:\мониторинг\Hujjatlar_Baza"

DISTRICTS_19 = [
    "Марғилон шаҳар", "Фарғона шаҳар", "Қувасой шаҳар", "Қўқон шаҳар",
    "Бағдод тумани", "Бешариқ тумани", "Бувайда тумани", "Данғара тумани",
    "Ёзёвон тумани", "Олтиариқ тумани", "Қўштепа тумани", "Риштон тумани",
    "Сўх тумани", "Тошлоқ тумани", "Учкўприк тумани", "Фарғона тумани",
    "Фурқат тумани", "Ўзбекистон тумани", "Қува тумани"
]

NOMINAL_SHEET_KEYWORDS = ["ф.и.ш", "рўйхат", "руйҳат", "фаблоси", "профилактик назорат", "умумий рўйх"]

NOMINAL_TO_SVOD_CONFIG = {
    "4-ички назорат Ф.И.Ш": {
        "svod_sheet": "2-ички назорат",
        "svod_col_idx": 3,
        "svod_title": "Ички назоратга олинганлар",
        "fish_col_idx": 3,
        "school_col_idx": 2,
        "has_section_header": True
    },
    "7-Эрта туғруқ Ф.И.Ш": {
        "svod_sheet": "6-эрта туғруқ",
        "svod_col_idx": 3,
        "svod_title": "Эрта туғруқ ҳолатлари",
        "fish_col_idx": 3,
        "school_col_idx": 2,
        "has_section_header": True
    },
    "9-қасд рўйхати": {
        "svod_sheet": "8-Ўз жонига қасд",
        "svod_col_idx": 3,
        "svod_title": "Ўз жонига қасд қилиш ҳолатлари",
        "fish_col_idx": 3,
        "school_col_idx": 2,
        "has_section_header": True
    },
    "11-жиноят фаблоси": {
        "svod_sheet": "10-жиноят сони",
        "svod_col_idx": 6,
        "svod_title": "Жиноят содир этган ўқувчилар",
        "fish_col_idx": 4,
        "school_col_idx": 3,
        "district_col_idx": 2,
        "has_section_header": False
    },
    "13-чет элга кетган. умумий рўйх": {
        "svod_sheet": "12-Чет элга кетган ",
        "svod_col_idx": 3,
        "svod_title": "Чет элга кетган ўқувчилар",
        "fish_col_idx": 4,
        "school_col_idx": 3,
        "district_col_idx": 2,
        "has_section_header": True
    },
    "14 Сурункали руйҳат": {
        "svod_sheet": "14 Сурункали дарс колдирилган",
        "svod_col_idx": 4,
        "svod_title": "Сурункали дарс қолдирганлар",
        "fish_col_idx": 4,
        "school_col_idx": 3,
        "district_col_idx": 2,
        "has_section_header": True
    }
}

def is_nominal_sheet(sheet_name: str) -> bool:
    s_lower = sheet_name.lower()
    return any(k in s_lower for k in NOMINAL_SHEET_KEYWORDS)

def get_sheet_columns(sheet_name: str):
    """
    Analyzes header rows (row_index <= 5) to produce unified column titles.
    """
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT row_index, cells_json, is_header 
    FROM excel_15_talik_sheets 
    WHERE sheet_name = ? AND row_index <= 5 
    ORDER BY row_index
    """, (sheet_name,))
    header_rows = cursor.fetchall()
    conn.close()

    if not header_rows:
        return []

    # Filter out title banner rows (rows with <= 2 non-empty cells or very short content)
    candidate_rows = []
    max_len = 0
    for r in header_rows:
        cells = json.loads(r['cells_json'])
        non_empty = [c for c in cells if str(c).strip()]
        if len(non_empty) >= 3 or any(x in " ".join(cells).lower() for x in ["т/р", "№", "ҳудудлар", "туман"]):
            candidate_rows.append(cells)
            if len(cells) > max_len:
                max_len = len(cells)

    if not candidate_rows:
        # Fallback to all rows
        for r in header_rows:
            cells = json.loads(r['cells_json'])
            candidate_rows.append(cells)
            if len(cells) > max_len:
                max_len = len(cells)

    # Pad all candidate rows to max_len
    for p in candidate_rows:
        while len(p) < max_len:
            p.append("")

    main_row = candidate_rows[0] if candidate_rows else [""] * max_len
    sub_row = candidate_rows[1] if len(candidate_rows) > 1 else [""] * max_len

    # Propagate merged parent headers across empty cells
    propagated_main = list(main_row)
    last_val = ""
    for c_i in range(len(propagated_main)):
        val = propagated_main[c_i].strip()
        if val:
            last_val = val
        elif c_i > 2 and last_val and (c_i < len(sub_row) and sub_row[c_i].strip()):
            propagated_main[c_i] = last_val

    columns = []
    for c_i in range(max_len):
        m_txt = propagated_main[c_i].strip().replace("\n", " ") if c_i < len(propagated_main) else ""
        s_txt = sub_row[c_i].strip().replace("\n", " ") if c_i < len(sub_row) else ""

        if m_txt and s_txt and m_txt != s_txt:
            full_title = f"{m_txt} ({s_txt})"
        elif s_txt:
            full_title = s_txt
        elif m_txt:
            full_title = m_txt
        else:
            full_title = f"Ustun {c_i}"

        # Clean multiple spaces
        full_title = " ".join(full_title.split())

        # Determine type & readonly status
        is_ro = False
        if c_i == 0 and (full_title == "Ustun 0" or not full_title):
            is_ro = True
        elif c_i == 1 and any(x in full_title.lower() for x in ["№", "т/р", "t/r"]):
            is_ro = True
        elif ("туман" in full_title.lower() or "ҳудуд" in full_title.lower() or "шаҳар" in full_title.lower()) and not is_nominal_sheet(sheet_name):
            is_ro = True # District name is fixed in svod sheets

        columns.append({
            "index": c_i,
            "title": full_title,
            "is_readonly": is_ro
        })

    return columns

def recalculate_svod_totals(conn, sheet_name: str):
    """
    Recalculates the 'Жами' / 'Вилоят бўйича жами' row by summing all district rows.
    """
    cursor = conn.cursor()
    cursor.execute("""
    SELECT row_index, cells_json, is_header, is_total 
    FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (sheet_name,))
    rows = cursor.fetchall()

    if not rows:
        return

    district_rows = []
    total_row = None

    for r in rows:
        r_idx = r['row_index'] if isinstance(r, (sqlite3.Row, dict)) else r[0]
        r_cells = r['cells_json'] if isinstance(r, (sqlite3.Row, dict)) else r[1]
        r_hdr = r['is_header'] if isinstance(r, (sqlite3.Row, dict)) else r[2]

        cells = json.loads(r_cells)
        cell_text = " ".join(str(x) for x in cells[:3]).lower()
        if ("вилоят бўйича" in cell_text or "бўйича жами" in cell_text or "вилоят  бўйича" in cell_text or cell_text.strip().endswith("жами:")) and not r_hdr:
            total_row = (r_idx, r_cells)
        elif not r_hdr and r_idx >= 5:
            # Check if this row is a district row
            row_str = " ".join(str(c) for c in cells[:4])
            if any(d in row_str for d in DISTRICTS_19):
                district_rows.append(cells)

    if not total_row or not district_rows:
        return

    tot_cells = json.loads(total_row[1])
    num_cols = len(tot_cells)

    # Sum numeric columns
    for c_i in range(num_cols):
        # Check if this column contains numbers in district rows
        is_num_col = False
        col_sum = 0
        has_at_least_one_number = False

        for d_cells in district_rows:
            if c_i < len(d_cells):
                raw = str(d_cells[c_i]).strip().replace(" ", "").replace(",", ".")
                try:
                    val = float(raw)
                    col_sum += val
                    has_at_least_one_number = True
                    is_num_col = True
                except ValueError:
                    pass

        # If it's the title cell (e.g. contains "Жами:"), keep it
        if "жами" in str(tot_cells[c_i]).lower():
            continue
        elif c_i in [0, 1, 2] and not has_at_least_one_number:
            continue

        if is_num_col and has_at_least_one_number:
            if col_sum.is_integer():
                tot_cells[c_i] = str(int(col_sum))
            else:
                tot_cells[c_i] = f"{col_sum:.1f}"

    cursor.execute("""
    UPDATE excel_15_talik_sheets 
    SET cells_json = ?, is_total = 1 
    WHERE sheet_name = ? AND row_index = ?
    """, (json.dumps(tot_cells, ensure_ascii=False), sheet_name, total_row[0]))
    conn.commit()

def save_svod_district_row(sheet_name: str, district_name: str, updated_cells: list, user_name: str = "Tuman mas'uli"):
    """
    Updates the district row in a svod table and recalculates totals.
    """
    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
    SELECT row_index, cells_json 
    FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (sheet_name,))
    rows = cursor.fetchall()

    target_row_index = None
    target_cells = None
    for r in rows:
        c = json.loads(r['cells_json'])
        # Match district name in any cell (usually cell 2 or 3)
        if any(district_name in str(cell) for cell in c):
            target_row_index = r['row_index']
            target_cells = c
            break

    if target_row_index is None:
        conn.close()
        return False, f"'{sheet_name}' varag'ida '{district_name}' qatori topilmadi."

    # Merge updated cells with existing cells
    final_cells = list(target_cells)
    while len(final_cells) < len(updated_cells):
        final_cells.append("")

    for i in range(len(updated_cells)):
        if i == 0 and not str(updated_cells[i]).strip():
            continue # Preserve padding
        val = str(updated_cells[i]).strip()
        final_cells[i] = val

    cursor.execute("""
    UPDATE excel_15_talik_sheets 
    SET cells_json = ? 
    WHERE sheet_name = ? AND row_index = ?
    """, (json.dumps(final_cells, ensure_ascii=False), sheet_name, target_row_index))
    conn.commit()

    # Recalculate totals
    recalculate_svod_totals(conn, sheet_name)
    conn.close()

    log_audit("15-talik Tahrirlandi", f"{user_name}: '{sheet_name}' varag'ida '{district_name}' ma'lumotlari kiritildi va saqlandi.")
    return True, f"'{district_name}' ko‘rsatkichlari muvaffaqiyatli saqlandi va jami hisob-kitoblar yangilandi!"

def normalize_student_fio(fio: str) -> str:
    if not fio:
        return ""
    s = str(fio).strip().lower()
    s = re.sub(r'[\'"`‘ʼ’«»]', '', s)
    s = re.sub(r'\d{2}[./-]\d{2}[./-]\d{4}.*$', '', s)
    s = re.sub(r'\d{4}\s*йил.*$', '', s)
    return " ".join(s.split())

def get_svod_target_count(district_name: str, nominal_sheet_name: str) -> dict:
    cfg = NOMINAL_TO_SVOD_CONFIG.get(nominal_sheet_name)
    if not cfg:
        return {"has_svod": False, "target_count": 0, "svod_sheet": "", "title": ""}
    
    svod_sheet = cfg["svod_sheet"]
    col_idx = cfg["svod_col_idx"]
    title = cfg["svod_title"]
    
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT cells_json FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (svod_sheet,))
    rows = cursor.fetchall()
    conn.close()

    target = 0
    dist_clean = district_name.replace("тумани", "").replace("шаҳар", "").replace("шаҳри", "").strip().lower()

    for r in rows:
        cells = json.loads(r['cells_json'])
        row_str = " ".join(str(x) for x in cells).lower()
        if dist_clean in row_str:
            if col_idx < len(cells):
                raw = str(cells[col_idx]).strip().replace(" ", "").replace(",", ".")
                try:
                    target = int(float(raw))
                except (ValueError, TypeError):
                    target = 0
            break

    return {
        "has_svod": True,
        "target_count": target,
        "svod_sheet": svod_sheet,
        "title": title
    }

def get_district_nominal_records(sheet_name: str, district_name: str) -> list:
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT row_index, cells_json FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (sheet_name,))
    all_rows = cursor.fetchall()
    conn.close()

    cfg = NOMINAL_TO_SVOD_CONFIG.get(sheet_name, {})
    fish_col = cfg.get("fish_col_idx", 3)
    has_sec = cfg.get("has_section_header", True)

    dist_clean = district_name.replace("тумани", "").replace("шаҳар", "").replace("шаҳри", "").strip().lower()

    records = []
    in_section = False

    for r in all_rows:
        r_idx = r['row_index']
        cells = json.loads(r['cells_json'])
        row_str = " ".join(str(x) for x in cells)

        if has_sec and (len(cells) <= 4 or len([x for x in cells if str(x).strip()]) <= 2):
            if any(d in row_str for d in ["тумани", "шаҳар", "шаҳри"]):
                if dist_clean in row_str.lower():
                    in_section = True
                    continue
                elif in_section:
                    in_section = False
                    break
        elif in_section:
            if any("жами" in str(x).lower() for x in cells):
                break
            fish_val = cells[fish_col].strip() if fish_col < len(cells) else ""
            if fish_val or (len(cells) > 2 and str(cells[1]).strip().isdigit()):
                records.append({
                    "row_index": r_idx,
                    "cells": cells,
                    "fish": fish_val,
                    "norm_fish": normalize_student_fio(fish_val)
                })
        elif not has_sec or not in_section:
            if dist_clean in row_str.lower() and r_idx > 3 and len(cells) > 3:
                fish_val = cells[fish_col].strip() if fish_col < len(cells) else ""
                records.append({
                    "row_index": r_idx,
                    "cells": cells,
                    "fish": fish_val,
                    "norm_fish": normalize_student_fio(fish_val)
                })

    return records

def generate_nominal_template(sheet_name: str, district_name: str = "") -> tuple:
    cols = get_sheet_columns(sheet_name)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    clean_title = re.sub(r'[\/\\?*:[\]]', '_', sheet_name)[:30]
    ws.title = clean_title
    
    export_cols = []
    for c in cols:
        if c['index'] == 0 and ("Ustun" in c['title'] or not c['title']):
            continue
        export_cols.append(c)

    title_font = Font(name="Arial", size=12, bold=True, color="1E3A8A")
    hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    dist_txt = f" ({district_name})" if district_name else ""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(export_cols))
    ws.cell(row=1, column=1, value=f"{sheet_name} bo‘yicha ma’lumotlarni kiritish namunaviy shabloni{dist_txt}").font = title_font
    ws.row_dimensions[1].height = 28
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[2].height = 36
    for col_idx, col_data in enumerate(export_cols, 1):
        cell = ws.cell(row=2, column=col_idx, value=col_data['title'])
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.border = border_thin
        cell.alignment = align_center

    ws.row_dimensions[3].height = 24
    sample_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    sample_font = Font(name="Arial", size=9, italic=True, color="166534")

    for col_idx, col_data in enumerate(export_cols, 1):
        t_low = col_data['title'].lower()
        val = ""
        if any(x in t_low for x in ["т/р", "№", "t/r"]):
            val = "1"
        elif "туман" in t_low or "шаҳар" in t_low:
            val = district_name or "Бувайда тумани"
        elif "мактаб" in t_low or "синф" in t_low:
            val = "1-мактаб 8-А синф"
        elif "ф.и.ш" in t_low or "исми" in t_low or "шарифи" in t_low:
            val = "Каримов Жасурбек Анвар ўғли"
        elif "сана" in t_low or "сабаб" in t_low:
            val = "15.09.2025 мактаб ички тартиб қоидасини бузган"
        elif "масъул" in t_low or "инспектор" in t_low:
            val = "А.Турсунов 945532402"
        elif "тўгарак" in t_low:
            val = "Футбол тўгараги"
        elif "манзил" in t_low:
            val = "Янгиобод МФЙ, Тинчлик кўчаси 12-уй"
        else:
            val = "Намунавий маълумот"

        cell = ws.cell(row=3, column=col_idx, value=val)
        cell.font = sample_font
        cell.fill = sample_fill
        cell.border = border_thin
        cell.alignment = align_center if col_idx <= 2 else align_left

    for col_idx, col_data in enumerate(export_cols, 1):
        col_letter = get_column_letter(col_idx)
        title_len = len(str(col_data['title']))
        ws.column_dimensions[col_letter].width = max(min(title_len + 4, 35), 14)

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()
    wb.close()
    
    filename = f"Shablon_{clean_title}_{district_name or 'Tuman'}.xlsx"
    return filename, file_bytes

def parse_excel_nominal_rows(file_bytes: bytes, sheet_name: str, district_name: str) -> tuple:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    cfg = NOMINAL_TO_SVOD_CONFIG.get(sheet_name, {})
    
    hdr_row_idx = 1
    found_hdr = False
    col_mapping = {}

    for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if r_idx > 8:
            break
        row_str = " ".join(str(c) for c in row if c is not None).lower()
        if any(k in row_str for k in ["ф.и.ш", "исми", "шарифи", "ўқувчи", "мактаб", "синф", "№", "т/р"]):
            hdr_row_idx = r_idx
            found_hdr = True
            for c_idx, cell_val in enumerate(row):
                if not cell_val:
                    continue
                cv_low = str(cell_val).lower()
                if any(x in cv_low for x in ["ф.и.ш", "исми", "шарифи", "ўқувчи"]):
                    col_mapping["fish"] = c_idx
                elif any(x in cv_low for x in ["мактаб", "синф"]):
                    col_mapping["school"] = c_idx
                elif any(x in cv_low for x in ["№", "т/р", "t/r"]):
                    col_mapping["tr"] = c_idx
            break

    if "fish" not in col_mapping:
        col_mapping["fish"] = 2

    raw_rows = []
    seen_in_file = set()
    duplicates_in_file = []

    for r_idx, row in enumerate(ws.iter_rows(min_row=hdr_row_idx + 1, values_only=True), hdr_row_idx + 1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if not any(cells):
            continue

        fish_idx = col_mapping.get("fish", 2)
        fish_val = cells[fish_idx] if fish_idx < len(cells) else ""
        
        # If fish_val is empty, search other cells for student name
        if not fish_val:
            for c_cand in cells:
                w = c_cand.split()
                if len(w) >= 2 and not any(ch.isdigit() for ch in c_cand) and not any(x in c_cand.lower() for x in ["мактаб", "синф", "туман"]):
                    fish_val = c_cand
                    break

        if not fish_val or any(k in fish_val.lower() for k in ["ф.и.ш", "намунавий", "каримов жасурбек", "фамилияси"]):
            continue

        norm_fish = normalize_student_fio(fish_val)
        if not norm_fish:
            continue

        if norm_fish in seen_in_file:
            duplicates_in_file.append({"row": r_idx, "fish": fish_val})
            continue

        seen_in_file.add(norm_fish)
        raw_rows.append({
            "excel_row": r_idx,
            "fish": fish_val,
            "norm_fish": norm_fish,
            "cells": cells
        })

    wb.close()
    return raw_rows, duplicates_in_file

def format_row_for_database(sheet_name: str, district_name: str, tr_num: int, input_cells: list) -> list:
    cfg = NOMINAL_TO_SVOD_CONFIG.get(sheet_name, {})
    has_sec = cfg.get("has_section_header", True)
    
    cells_data = list(input_cells)
    if cells_data and str(cells_data[0]).strip().isdigit():
        cells_data = cells_data[1:]
        
    final_row = ["", str(tr_num)]
    
    if not has_sec or sheet_name in ["11-жиноят фаблоси", "13-чет элга кетган. умумий рўйх", "14 Сурункали руйҳат"]:
        final_row.append(district_name)
        
    for val in cells_data:
        if val and district_name.lower() in str(val).lower() and len(final_row) == 2:
            continue
        final_row.append(str(val).strip())
        
    return final_row

def import_nominal_excel(file_bytes: bytes, sheet_name: str, district_name: str, mode: str = "append", user_name: str = "Tuman mas'uli") -> dict:
    conn = get_db()
    cursor = conn.cursor()

    valid_rows, file_dups = parse_excel_nominal_rows(file_bytes, sheet_name, district_name)
    if not valid_rows:
        conn.close()
        return {
            "status": "error",
            "message": "Faylda yaroqli o‘quvchi ma’lumotlari topilmadi. Iltimos namunaviy shablon bo‘yicha to‘ldirilganini tekshiring."
        }

    existing_records = get_district_nominal_records(sheet_name, district_name)
    existing_norm_fios = {r['norm_fish'] for r in existing_records if r['norm_fish']}

    filtered_rows = []
    db_dups = []
    for r in valid_rows:
        if mode == "append" and r['norm_fish'] in existing_norm_fios:
            db_dups.append({"fish": r['fish'], "excel_row": r['excel_row']})
        else:
            filtered_rows.append(r)

    if not filtered_rows:
        conn.close()
        dup_names = ", ".join([d['fish'] for d in (file_dups + db_dups)[:5]])
        return {
            "status": "error",
            "message": f"Barcha o‘quvchilar allaqachon ro‘yxatda mavjud yoki faylda takrorlangan: {dup_names}"
        }

    svod_info = get_svod_target_count(district_name, sheet_name)
    target_count = svod_info.get("target_count", 0)

    prev_count = len(existing_records)
    new_added = len(filtered_rows)
    final_count = new_added if mode == "replace" else (prev_count + new_added)
    diff = final_count - target_count

    cursor.execute("""
    SELECT id, row_index, cells_json, sheet_index 
    FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (sheet_name,))
    all_rows = cursor.fetchall()
    if not all_rows:
        conn.close()
        return {"status": "error", "message": "Varaq ma'lumotlari topilmadi"}

    sheet_index = all_rows[0]['sheet_index']
    cfg = NOMINAL_TO_SVOD_CONFIG.get(sheet_name, {})
    has_sec = cfg.get("has_section_header", True)
    dist_clean = district_name.replace("тумани", "").replace("шаҳар", "").replace("шаҳри", "").strip().lower()

    if mode == "replace" and existing_records:
        old_indexes = [r['row_index'] for r in existing_records]
        placeholders = ",".join("?" * len(old_indexes))
        cursor.execute(f"""
        DELETE FROM excel_15_talik_sheets 
        WHERE sheet_name = ? AND row_index IN ({placeholders})
        """, [sheet_name] + old_indexes)

    cursor.execute("""
    SELECT row_index, cells_json 
    FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (sheet_name,))
    remaining_rows = cursor.fetchall()

    insert_after_idx = None
    section_hdr_idx = None
    start_tr = 1 if mode == "replace" else (prev_count + 1)

    for r in remaining_rows:
        c = json.loads(r['cells_json'])
        r_str = " ".join(str(x) for x in c)
        if has_sec and (len(c) <= 4 or len([x for x in c if str(x).strip()]) <= 2):
            if any(d in r_str for d in ["тумани", "шаҳар", "шаҳри"]):
                if dist_clean in r_str.lower():
                    section_hdr_idx = r['row_index']
                    insert_after_idx = r['row_index']
                elif section_hdr_idx is not None:
                    break
        elif section_hdr_idx is not None:
            if any("жами" in str(x).lower() for x in c):
                break
            insert_after_idx = r['row_index']
        elif not has_sec and dist_clean in r_str.lower():
            insert_after_idx = r['row_index']

    if has_sec and section_hdr_idx is None:
        last_idx = remaining_rows[-1]['row_index'] if remaining_rows else 10
        section_hdr_idx = last_idx + 1
        sec_cells = ["", district_name]
        cursor.execute("""
        INSERT INTO excel_15_talik_sheets (sheet_index, sheet_name, row_index, cells_json, is_header, is_total)
        VALUES (?, ?, ?, ?, 0, 0)
        """, (sheet_index, sheet_name, section_hdr_idx, json.dumps(sec_cells, ensure_ascii=False)))
        insert_after_idx = section_hdr_idx

    if insert_after_idx is None:
        insert_after_idx = remaining_rows[-1]['row_index'] if remaining_rows else 10

    for i, r_data in enumerate(filtered_rows):
        cur_tr = start_tr + i
        formatted_cells = format_row_for_database(sheet_name, district_name, cur_tr, r_data['cells'])
        temp_idx = insert_after_idx + 0.001 * (i + 1)
        cursor.execute("""
        INSERT INTO excel_15_talik_sheets (sheet_index, sheet_name, row_index, cells_json, is_header, is_total)
        VALUES (?, ?, ?, ?, 0, 0)
        """, (sheet_index, sheet_name, temp_idx, json.dumps(formatted_cells, ensure_ascii=False)))

    cursor.execute("""
    SELECT id FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index ASC, id ASC
    """, (sheet_name,))
    all_sheet_ids = [row[0] for row in cursor.fetchall()]

    for new_idx, record_id in enumerate(all_sheet_ids, 1):
        cursor.execute("""
        UPDATE excel_15_talik_sheets 
        SET row_index = ? 
        WHERE id = ?
        """, (new_idx, record_id))

    conn.commit()
    conn.close()

    log_audit("15-talik Excel Import", f"{user_name}: '{sheet_name}' varag'iga '{district_name}' bo'yicha {new_added} ta o'quvchi import qilindi. ({mode} rejimi)")

    return {
        "status": "success",
        "mode": mode,
        "imported_count": new_added,
        "target_count": target_count,
        "final_count": final_count,
        "diff": diff,
        "is_matched": (diff == 0 if target_count > 0 else True),
        "duplicates_in_file": file_dups,
        "duplicates_in_db": db_dups,
        "message": f"Exceldan {new_added} ta o‘quvchi muvaffaqiyatli yuklandi!"
    }

def add_nominal_record(sheet_name: str, district_name: str, row_data: list, user_name: str = "Tuman mas'uli"):
    """
    Adds a new nominal record (e.g. student in 4-ички назорат or incident in 11-фаблоси)
    under the district's section with duplicate prevention.
    """
    # Check duplicate student
    existing_recs = get_district_nominal_records(sheet_name, district_name)
    existing_norms = {r['norm_fish'] for r in existing_recs if r['norm_fish']}
    
    incoming_fish = ""
    for item in row_data:
        val = str(item).strip()
        words = val.split()
        if len(words) >= 2 and not any(ch.isdigit() for ch in val) and not any(k in val.lower() for k in ["мактаб", "синф", "мфй"]):
            incoming_fish = val
            break
    if not incoming_fish and len(row_data) >= 2:
        incoming_fish = str(row_data[1] if len(row_data) > 1 else row_data[0]).strip()

    if incoming_fish:
        norm_in = normalize_student_fio(incoming_fish)
        if norm_in and norm_in in existing_norms:
            return False, f"Diqqat: '{incoming_fish}' mazkur ro‘yxatda allaqachon mavjud! Bir xil o‘quvchini takroran kiritish taqiqlanadi."

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
    SELECT id, row_index, cells_json, sheet_index 
    FROM excel_15_talik_sheets 
    WHERE sheet_name = ? 
    ORDER BY row_index
    """, (sheet_name,))
    all_rows = cursor.fetchall()

    if not all_rows:
        conn.close()
        return False, "Varaq ma'lumotlari topilmadi"

    sheet_index = all_rows[0]['sheet_index']

    # Locate district section or where to insert
    # In nominal sheets, district sections are either:
    # 1) Row like ['', 'Бувайда тумани']
    # 2) Or each row has district in column 2 / 3
    section_hdr_index = None
    insert_after_row_index = None
    district_count = 0

    for idx, r in enumerate(all_rows):
        c = json.loads(r['cells_json'])
        # Section header check: small row with just district name
        if len(c) <= 4 and any(district_name.lower() in str(cell).lower() for cell in c):
            section_hdr_index = r['row_index']
            insert_after_row_index = r['row_index']
            # Find the last record under this section
            for next_r in all_rows[idx + 1:]:
                next_c = json.loads(next_r['cells_json'])
                # If next_c is another district section or total, stop
                if len(next_c) <= 4 and any(d.lower() in " ".join(str(x) for x in next_c).lower() for d in DISTRICTS_19):
                    break
                if any("жами" in str(x).lower() for x in next_c):
                    break
                insert_after_row_index = next_r['row_index']
                district_count += 1
            break
        # Or individual rows check
        elif any(district_name.lower() in str(cell).lower() for cell in c):
            insert_after_row_index = r['row_index']
            district_count += 1

    # If neither found, insert at the end before total row
    if insert_after_row_index is None:
        last_row = all_rows[-1]
        insert_after_row_index = last_row['row_index']

        # Add section header for this district if not present
        new_sec_row_index = insert_after_row_index + 1
        sec_cells = ["", district_name]
        cursor.execute("""
        INSERT INTO excel_15_talik_sheets (sheet_index, sheet_name, row_index, cells_json, is_header, is_total)
        VALUES (?, ?, ?, ?, 0, 0)
        """, (sheet_index, sheet_name, new_sec_row_index, json.dumps(sec_cells, ensure_ascii=False)))
        insert_after_row_index = new_sec_row_index

    new_row_index = insert_after_row_index + 1
    new_num = district_count + 1

    # Prepare cells
    # Typically: cell 0: '', cell 1: number (e.g. str(new_num)), then inputs
    cleaned_row = ["", str(new_num)]
    # If the sheet format has district in cell 2, ensure district is present
    cols = get_sheet_columns(sheet_name)
    has_district_col = any("туман" in c['title'].lower() or "ҳудуд" in c['title'].lower() for c in cols if c['index'] == 2)
    
    start_col = 2
    if has_district_col and len(row_data) > 0 and district_name not in str(row_data[0]):
        cleaned_row.append(district_name)
        start_col = 3

    for item in row_data:
        cleaned_row.append(str(item).strip())

    # Shift subsequent row_indexes by 1 to make room
    cursor.execute("""
    UPDATE excel_15_talik_sheets 
    SET row_index = row_index + 1 
    WHERE sheet_name = ? AND row_index >= ?
    """, (sheet_name, new_row_index))

    # Insert new row
    cursor.execute("""
    INSERT INTO excel_15_talik_sheets (sheet_index, sheet_name, row_index, cells_json, is_header, is_total)
    VALUES (?, ?, ?, ?, 0, 0)
    """, (sheet_index, sheet_name, new_row_index, json.dumps(cleaned_row, ensure_ascii=False)))

    conn.commit()
    conn.close()

    log_audit("15-talik Ro'yxat Qo'shildi", f"{user_name}: '{sheet_name}' varag'iga '{district_name}' bo'yicha yangi ma'lumot ({row_data[0] if row_data else ''}) qo'shildi.")
    return True, f"Yangi ma'lumot '{district_name}' bo'limiga muvaffaqiyatli qo'shildi!"

def update_nominal_record(sheet_name: str, row_index: int, row_data: list, user_name: str = "Tuman mas'uli"):
    """
    Updates an existing nominal record.
    """
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT cells_json FROM excel_15_talik_sheets 
    WHERE sheet_name = ? AND row_index = ?
    """, (sheet_name, row_index))
    r = cursor.fetchone()
    if not r:
        conn.close()
        return False, "Qator topilmadi"

    orig_cells = json.loads(r['cells_json'])
    final_cells = list(orig_cells)

    while len(final_cells) < len(row_data):
        final_cells.append("")

    for i in range(len(row_data)):
        # Preserve index 0 and 1 if row_data doesn't override
        if i < 2 and not str(row_data[i]).strip():
            continue
        final_cells[i] = str(row_data[i]).strip()

    cursor.execute("""
    UPDATE excel_15_talik_sheets 
    SET cells_json = ? 
    WHERE sheet_name = ? AND row_index = ?
    """, (json.dumps(final_cells, ensure_ascii=False), sheet_name, row_index))
    conn.commit()
    conn.close()

    log_audit("15-talik Tahrirlandi", f"{user_name}: '{sheet_name}' varag'ida {row_index}-qator ma'lumotlari yangilandi.")
    return True, "Ma'lumot muvaffaqiyatli yangilandi!"

def delete_nominal_record(sheet_name: str, row_index: int, user_name: str = "Tuman mas'uli"):
    """
    Deletes a nominal record row.
    """
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT cells_json FROM excel_15_talik_sheets 
    WHERE sheet_name = ? AND row_index = ?
    """, (sheet_name, row_index))
    r = cursor.fetchone()
    if not r:
        conn.close()
        return False, "Qator topilmadi"

    c = json.loads(r['cells_json'])
    sample_text = " ".join(c[:4])

    cursor.execute("""
    DELETE FROM excel_15_talik_sheets 
    WHERE sheet_name = ? AND row_index = ?
    """, (sheet_name, row_index))

    # Shift back
    cursor.execute("""
    UPDATE excel_15_talik_sheets 
    SET row_index = row_index - 1 
    WHERE sheet_name = ? AND row_index > ?
    """, (sheet_name, row_index))

    conn.commit()
    conn.close()

    log_audit("15-talik O'chirildi", f"{user_name}: '{sheet_name}' varag'idan qator o'chirildi: {sample_text}")
    return True, "Qator muvaffaqiyatli o'chirildi!"

def export_updated_15_talik_workbook(month_num: int = 8, year: int = 2026):
    """
    Takes the master Excel workbook as a layout/styling template,
    injects all current rows from `excel_15_talik_sheets` SQLite table,
    and produces the fresh updated Excel workbook.
    Also saves a backup in the user's permanent folder.
    """
    MONTH_NAMES_UZ = {
        1: "Yanvar", 2: "Fevral", 3: "Mart", 4: "Aprel",
        5: "May", 6: "Iyun", 7: "Iyul", 8: "Avgust",
        9: "Sentabr", 10: "Oktabr", 11: "Noyabr", 12: "Dekabr"
    }
    m_uz = MONTH_NAMES_UZ.get(int(month_num), "Avgust")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"15_Talik_Rasmiy_Kitob_{m_uz}_{year}_{timestamp}.xlsx"
    out_path = os.path.join(EXPORTS_DIR, out_name)

    if not os.path.exists(MASTER_EXCEL_PATH):
        # Fallback empty workbook
        wb = openpyxl.Workbook()
        wb.save(out_path)
        return out_name, out_path

    # Load master template to retain all fonts, borders, column widths and merged cells
    wb = openpyxl.load_workbook(MASTER_EXCEL_PATH)

    conn = get_db()
    cursor = conn.cursor()

    for sname in wb.sheetnames:
        sheet = wb[sname]
        cursor.execute("""
        SELECT row_index, cells_json, is_header, is_total 
        FROM excel_15_talik_sheets 
        WHERE sheet_name = ? 
        ORDER BY row_index
        """, (sname,))
        db_rows = cursor.fetchall()

        if not db_rows:
            continue

        for r_data in db_rows:
            r_idx = r_data['row_index']
            cells = json.loads(r_data['cells_json'])
            is_hdr = r_data['is_header']

            # If it's not a header row, write database cell values
            # (Header rows can be preserved or updated if matching)
            for c_idx, val in enumerate(cells, 1):
                if val is None or val == "":
                    continue

                cell_obj = sheet.cell(row=r_idx, column=c_idx)

                # Try numeric conversion if it looks like an integer/float
                if not is_hdr and str(val).isdigit():
                    cell_obj.value = int(val)
                else:
                    try:
                        if not is_hdr and "." in str(val) and float(val):
                            cell_obj.value = float(val)
                        else:
                            cell_obj.value = str(val)
                    except ValueError:
                        cell_obj.value = str(val)

    conn.close()

    # Save to exports directory
    os.makedirs(EXPORTS_DIR, exist_ok=True)
    wb.save(out_path)

    # Also automatically save into user's permanent local storage
    # E:\мониторинг\Hujjatlar_Baza\15_Talik_Monitoring\...
    backup_dir = os.path.join(BASE_STORAGE_DIR, "15_Talik_Monitoring", f"{year}_{m_uz}")
    try:
        os.makedirs(backup_dir, exist_ok=True)
        permanent_backup = os.path.join(backup_dir, f"15_Talik_Viloyat_Kitobi_{m_uz}_{year}.xlsx")
        shutil.copy2(out_path, permanent_backup)
    except Exception as e:
        print(f"Warning: Could not save backup to {backup_dir}: {e}")

    wb.close()
    return out_name, out_path
