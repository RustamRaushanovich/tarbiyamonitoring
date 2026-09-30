# -*- coding: utf-8 -*-
import openpyxl
import os
import re
from app.db import get_db
from app.services.ai_service import analyze_incident_ai
from app.parsers.fabula_analyzer import parse_incident_fabula

def _normalize_str(val):
    if not val:
        return ""
    # Strip birth date in brackets if any and lower
    clean = re.sub(r'\(.*?\)', '', str(val))
    clean = re.sub(r'\d{2}\.\d{2}\.\d{4}', '', clean)
    clean = re.sub(r'[\s\.\,\-]+', ' ', clean).strip().lower()
    return clean

def import_monthly_incidents_file(filepath, auto_ai_analyze=True):
    """
    Parses a monthly juvenile crime Excel file.
    Compares against existing records:
    - Identifies NEW incidents (is_new = 1, days_remaining = 10)
    - Leaves existing incidents (is_new = 0)
    - If auto_ai_analyze=True, immediately generates Gemini AI root cause analysis and 10-day demand letter!
    """
    if not os.path.exists(filepath):
        return {"status": "error", "message": f"Fayl topilmadi: {filepath}"}

    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    conn = get_db()
    cursor = conn.cursor()

    # Get districts map
    cursor.execute("SELECT id, name FROM districts")
    districts = cursor.fetchall()
    districts_map = {d['name'].lower(): d['id'] for d in districts}

    # Fetch existing incidents for smart comparison
    cursor.execute("SELECT id, student_fio, district_name, criminal_article FROM juvenile_incidents")
    existing_rows = cursor.fetchall()
    existing_keys = set()
    for row in existing_rows:
        key = (_normalize_str(row['student_fio']), _normalize_str(row['district_name']))
        existing_keys.add(key)

    cursor.execute("SELECT COALESCE(MAX(case_number), 0) as max_case FROM juvenile_incidents")
    current_case_num = cursor.fetchone()['max_case']

    new_count = 0
    existing_count = 0
    new_incidents_list = []

    for r in range(5, sheet.max_row + 1):
        col3_tuman = sheet.cell(r, 3).value
        col4_maktab = sheet.cell(r, 4).value
        col5_jins = sheet.cell(r, 5).value
        col6_fish = sheet.cell(r, 6).value
        col7_manzil = sheet.cell(r, 7).value
        col8_tafsilot = sheet.cell(r, 8).value
        col9_modda = sheet.cell(r, 9).value
        col10_maktab_masul = sheet.cell(r, 10).value
        col11_mfy_masul = sheet.cell(r, 11).value
        col12_nozir = sheet.cell(r, 12).value
        col13_holat = sheet.cell(r, 13).value

        if not col6_fish or not str(col6_fish).strip():
            continue

        student_fio = str(col6_fish).strip()
        district_name = str(col3_tuman).strip() if col3_tuman else "Noma’lum"
        school_name = str(col4_maktab).strip() if col4_maktab else ""
        gender = 'Erkak' if str(col5_jins).strip().lower() in ['м', 'm', 'ўғил', 'эркак'] else 'Ayol'
        address = str(col7_manzil).strip() if col7_manzil else ""
        details = str(col8_tafsilot).strip() if col8_tafsilot else ""
        article = str(col9_modda).strip() if col9_modda else ""

        # Find district_id
        district_id = None
        for dname, did in districts_map.items():
            clean_d = dname.replace('шаҳар', '').replace('тумани', '').strip()
            if clean_d in district_name.lower():
                district_id = did
                break

        # Extract birth date
        bdate = None
        bdate_match = re.search(r'(\d{2}\.\d{2}\.\d{4})', student_fio)
        if bdate_match:
            bdate = bdate_match.group(1)

        key = (_normalize_str(student_fio), _normalize_str(district_name))
        is_already_present = key in existing_keys

        if is_already_present:
            existing_count += 1
            # Update existing action data if provided in newer file
            cursor.execute("""
            UPDATE incident_actions
            SET school_responsible_fio = COALESCE(NULLIF(?, ''), school_responsible_fio),
                mfy_responsible_fio = COALESCE(NULLIF(?, ''), mfy_responsible_fio),
                inspector_fio = COALESCE(NULLIF(?, ''), inspector_fio),
                current_study_status = COALESCE(NULLIF(?, ''), current_study_status)
            WHERE incident_id IN (
                SELECT id FROM juvenile_incidents 
                WHERE student_fio = ? AND district_name = ?
            )
            """, (
                str(col10_maktab_masul).strip() if col10_maktab_masul else "",
                str(col11_mfy_masul).strip() if col11_mfy_masul else "",
                str(col12_nozir).strip() if col12_nozir else "",
                str(col13_holat).strip() if col13_holat else "",
                student_fio, district_name
            ))
        else:
            # BRAND NEW INCIDENT!
            current_case_num += 1
            new_count += 1
            existing_keys.add(key)

            # Parse fabula for school-time (Mon-Sat 08:00-14:00)
            fabula_meta = parse_incident_fabula(details)

            # Insert as new incident with 10-day active countdown
            cursor.execute("""
            INSERT INTO juvenile_incidents (
                case_number, district_id, district_name, school_name,
                student_fio, birth_date, gender, home_address,
                incident_details, criminal_article, incident_date, days_remaining, 
                status, is_new, incident_exact_time, incident_weekday, is_school_time, school_time_badge
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, date('now'), 10, 'Nazoratda', 1, ?, ?, ?, ?)
            """, (
                current_case_num, district_id, district_name, school_name,
                student_fio, bdate, gender, address,
                details, article,
                fabula_meta['exact_time'], fabula_meta['weekday'],
                fabula_meta['is_school_time'], fabula_meta['badge']
            ))
            incident_id = cursor.lastrowid

            # Insert initial 10-day action row
            cursor.execute("""
            INSERT INTO incident_actions (
                incident_id, school_responsible_fio, mfy_responsible_fio,
                inspector_fio, current_study_status, is_completed_in_10_days
            ) VALUES (?, ?, ?, ?, ?, 0)
            """, (
                incident_id,
                str(col10_maktab_masul).strip() if col10_maktab_masul else "",
                str(col11_mfy_masul).strip() if col11_mfy_masul else "",
                str(col12_nozir).strip() if col12_nozir else "",
                str(col13_holat).strip() if col13_holat else ""
            ))

            # Auto AI Analysis with Gemini if requested
            ai_data = None
            if auto_ai_analyze and details:
                try:
                    ai_data = analyze_incident_ai(
                        student_fio=student_fio,
                        district_name=district_name,
                        school_name=school_name,
                        article=article,
                        details=details,
                        birth_date=bdate or ""
                    )
                    cursor.execute("""
                    UPDATE juvenile_incidents
                    SET ai_analysis = ?, ai_risk_level = ?, ai_recommendations = ?, ai_letter_draft = ?
                    WHERE id = ?
                    """, (
                        ai_data.get("sabab", ""),
                        ai_data.get("xavf_darajasi", "O‘rta"),
                        json.dumps(ai_data.get("tavsiyalar", []), ensure_ascii=False),
                        ai_data.get("talabnoma_xati", ""),
                        incident_id
                    ))
                except Exception as e:
                    print(f"AI analysis failed for incident {incident_id}: {e}")

            new_incidents_list.append({
                "id": incident_id,
                "student_fio": student_fio,
                "district_name": district_name,
                "school_name": school_name,
                "criminal_article": article,
                "ai_risk": ai_data.get("xavf_darajasi", "O‘rta") if ai_data else "Tahlil kutilmoqda"
            })

    conn.commit()
    conn.close()

    return {
        "status": "success",
        "total_rows": new_count + existing_count,
        "new_count": new_count,
        "existing_count": existing_count,
        "new_incidents": new_incidents_list
    }

def parse_and_import_incidents(filepath=r'e:\мониторинг\2026 йил (август ойлик вирусиз).xlsx'):
    return import_monthly_incidents_file(filepath, auto_ai_analyze=False)

if __name__ == '__main__':
    parse_and_import_incidents()
