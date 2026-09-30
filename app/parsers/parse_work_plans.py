# -*- coding: utf-8 -*-
import docx
import os
import re
from datetime import datetime
from app.db import get_db

DISTRICT_NAMES = [
    'Марғилон', 'Фарғона', 'Қувасой', 'Қўқон', 'Бағдод', 'Бешариқ', 
    'Бувайда', 'Данғара', 'Ёзёвон', 'Олтиариқ', 'Қўштепа', 'Риштон', 
    'Сўх', 'Тошлоқ', 'Учкўприк', 'Фурқат', 'Ўзбекистон', 'Қува'
]

SECTION_TITLES = {
    1: 'Tashkiliy va nazorat faoliyati',
    2: 'Normativ-huquqiy hujjatlar va buyruqlar ijrosi, o‘rganishlar',
    3: 'Hay’at va apparat yig‘ilishi materiallari tayyorlash',
    4: 'Ko‘rik-tanlovlar, festivallar va musobaqalar tashkil etish',
    5: 'O‘quv-amaliy seminarlar, mahorat darslari va konferensiyalar'
}

def detect_district_in_text(text):
    for d in DISTRICT_NAMES:
        if d.lower() in text.lower():
            return d
    return None

def detect_school_in_text(text):
    match = re.search(r'(\d+)[- ]?(сонли|сон)?\s*(мактаб|umumiy|maktab)', text, re.IGNORECASE)
    if match:
        return match.group(1)
    return None

def parse_and_import_work_plan(filepath, owner_fio, owner_role, user_id=None, replace_existing=True):
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return False

    doc = docx.Document(filepath)
    conn = get_db()
    cursor = conn.cursor()

    if replace_existing:
        cursor.execute("UPDATE work_plans SET status = 'Arxiv' WHERE owner_fio = ? AND status = 'Tasdiqlangan'", (owner_fio,))

    title = os.path.basename(filepath).replace('.docx', '')
    cursor.execute("""
    INSERT INTO work_plans (user_id, owner_fio, owner_role, year, title, source_filename, status)
    VALUES (?, ?, ?, 2026, ?, ?, 'Tasdiqlangan')
    """, (user_id, owner_fio, owner_role, title, os.path.basename(filepath)))
    plan_id = cursor.lastrowid

    # Get districts map
    cursor.execute("SELECT id, name FROM districts")
    districts_map = {row['name']: row['id'] for row in cursor.fetchall()}

    item_count = 0
    assigned_count = 0

    # If document has multiple tables (e.g. Azimov, Turdiyev)
    if len(doc.tables) >= 5:
        for sec_idx, table in enumerate(doc.tables[:5], start=1):
            sec_title = SECTION_TITLES.get(sec_idx, f"{sec_idx}-bo‘lim")
            for r_idx, row in enumerate(table.rows[1:], start=1):
                cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                if not any(cells):
                    continue
                
                code = cells[0] if len(cells) > 0 else f"{sec_idx}.{r_idx}"
                desc = cells[1] if len(cells) > 1 else ""
                deadline = cells[2] if len(cells) > 2 else ""
                resp = cells[3] if len(cells) > 3 else owner_fio
                report = cells[4] if len(cells) > 4 else "Tahliliy ma’lumotnoma / Ijro hujjati"

                if not desc:
                    continue

                # Check territorial assignment
                target_district_id = None
                target_type = 'internal'
                detected_district = detect_district_in_text(desc)
                if detected_district:
                    for d_name, d_id in districts_map.items():
                        if detected_district.lower() in d_name.lower():
                            target_district_id = d_id
                            target_type = 'district'
                            break

                cursor.execute("""
                INSERT INTO work_plan_items (
                    plan_id, section_number, section_title, item_code,
                    task_description, deadline_text, report_form,
                    responsible_persons, target_type, target_district_id, status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Kutilmoqda')
                """, (plan_id, sec_idx, sec_title, code, desc, deadline, report, resp, target_type, target_district_id))
                item_id = cursor.lastrowid
                item_count += 1

                # If district was detected, create task assignment
                if target_district_id:
                    cursor.execute("""
                    INSERT INTO task_district_assignments (
                        work_plan_item_id, district_id, assigned_by_user_id,
                        assignment_title, assignment_instructions, deadline_date, status
                    ) VALUES (?, ?, ?, ?, ?, ?, 'Yangi')
                    """, (item_id, target_district_id, user_id, f"{code}: {desc[:100]}...", desc, deadline))
                    assigned_count += 1

    # Single table case (e.g. D. A'zamov)
    elif len(doc.tables) >= 1:
        table = doc.tables[0]
        for r_idx, row in enumerate(table.rows[1:], start=1):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            if not any(cells):
                continue

            code = cells[0] if len(cells) > 0 else str(r_idx)
            desc = cells[1] if len(cells) > 1 else ""
            deadline = cells[2] if len(cells) > 2 else ""
            report = cells[3] if len(cells) > 3 else "Tahliliy ma’lumot"

            if not desc:
                continue

            target_district_id = None
            target_type = 'internal'
            detected_district = detect_district_in_text(desc)
            if detected_district:
                for d_name, d_id in districts_map.items():
                    if detected_district.lower() in d_name.lower():
                        target_district_id = d_id
                        target_type = 'district'
                        break

            cursor.execute("""
            INSERT INTO work_plan_items (
                plan_id, section_number, section_title, item_code,
                task_description, deadline_text, report_form,
                responsible_persons, target_type, target_district_id, status
            ) VALUES (?, 1, 'Asosiy vazifalar', ?, ?, ?, ?, ?, ?, ?, 'Kutilmoqda')
            """, (plan_id, code, desc, deadline, report, owner_fio, target_type, target_district_id))
            item_id = cursor.lastrowid
            item_count += 1

            if target_district_id:
                cursor.execute("""
                INSERT INTO task_district_assignments (
                    work_plan_item_id, district_id, assigned_by_user_id,
                    assignment_title, assignment_instructions, deadline_date, status
                ) VALUES (?, ?, ?, ?, ?, ?, 'Yangi')
                """, (item_id, target_district_id, user_id, f"{code}: {desc[:100]}...", desc, deadline))
                assigned_count += 1

    conn.commit()
    conn.close()
    print(f"Imported {title}: {item_count} tasks, {assigned_count} territorial assignments.")
    return True

def import_all_work_plans(source_dir=r'e:\мониторинг'):
    plans = [
        ('Ish_rejasi_A.Azimov_2026 2.docx', 'A.Azimov', 'Metodist-kurator', 3),
        ('Ish_rejasi_R_Turdiyev_2026.docx', 'R.Turdiyev', 'Bo‘lim boshlig‘i', 4),
        ('Ish_rejasi_Д_Аъзамов_2026.docx', 'D.A’zamov', 'Yetakchi mutaxassis', 5)
    ]
    for filename, fio, role, uid in plans:
        fpath = os.path.join(source_dir, filename)
        if os.path.exists(fpath):
            parse_and_import_work_plan(fpath, fio, role, uid)

if __name__ == '__main__':
    import_all_work_plans()
