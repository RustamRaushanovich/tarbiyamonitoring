# -*- coding: utf-8 -*-
import openpyxl
import os
import sys
import json
import time

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(os.path.dirname(CURRENT_DIR))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from app.db import get_db

MASTER_EXCEL_PATH = r'e:\мониторинг\Фарғона_вилояти_15_талик_2026_йил_8_ойи_2_1.xlsx'

def init_15_talik_exact_cache():
    if not os.path.exists(MASTER_EXCEL_PATH):
        print("Master 15-talik excel not found:", MASTER_EXCEL_PATH)
        return

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS excel_15_talik_sheets (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sheet_index INTEGER,
        sheet_name TEXT,
        row_index INTEGER,
        cells_json TEXT,
        is_header INTEGER DEFAULT 0,
        is_total INTEGER DEFAULT 0
    )
    """)

    cursor.execute("SELECT COUNT(*) as cnt FROM excel_15_talik_sheets")
    existing_cnt = cursor.fetchone()['cnt']
    if existing_cnt > 100:
        print(f"excel_15_talik_sheets already cached ({existing_cnt} rows).")
        conn.close()
        return

    print("Caching all 18 sheets from master Excel into SQLite...")
    t0 = time.time()
    wb = openpyxl.load_workbook(MASTER_EXCEL_PATH, data_only=True)

    total_inserted = 0
    for s_idx, sname in enumerate(wb.sheetnames, 1):
        sheet = wb[sname]
        max_rows = min(sheet.max_row, 300) # up to 300 rows for large sheets
        max_cols = min(sheet.max_column, 65)

        for r in range(1, max_rows + 1):
            row_vals = [sheet.cell(r, c).value for c in range(1, max_cols + 1)]
            # Clean trailing Nones
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            if not any(v is not None for v in row_vals):
                continue

            cleaned = []
            for v in row_vals:
                if v is None:
                    cleaned.append("")
                elif isinstance(v, float) and v.is_integer():
                    cleaned.append(str(int(v)))
                else:
                    cleaned.append(str(v).strip())

            is_header = 1 if r <= 5 and any("№" in str(x) or "Т/р" in str(x) or "номи" in str(x) or "йил" in str(x) for x in cleaned) else 0
            is_total = 1 if any("жами" in str(x).lower() for x in cleaned) else 0

            cursor.execute("""
            INSERT INTO excel_15_talik_sheets (sheet_index, sheet_name, row_index, cells_json, is_header, is_total)
            VALUES (?, ?, ?, ?, ?, ?)
            """, (s_idx, sname, r, json.dumps(cleaned, ensure_ascii=False), is_header, is_total))
            total_inserted += 1

    conn.commit()
    conn.close()
    wb.close()
    print(f"Cached {total_inserted} rows across 18 sheets in {round(time.time() - t0, 2)} seconds.")

def reload_15_talik_from_file(excel_path: str):
    if not os.path.exists(excel_path):
        return False, "Fayl topilmadi"
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM excel_15_talik_sheets")
    conn.commit()

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    total_inserted = 0
    for s_idx, sname in enumerate(wb.sheetnames, 1):
        sheet = wb[sname]
        max_rows = min(sheet.max_row, 300)
        max_cols = min(sheet.max_column, 65)

        for r in range(1, max_rows + 1):
            row_vals = [sheet.cell(r, c).value for c in range(1, max_cols + 1)]
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            if not any(v is not None for v in row_vals):
                continue

            cleaned = []
            for v in row_vals:
                if v is None:
                    cleaned.append("")
                elif isinstance(v, float) and v.is_integer():
                    cleaned.append(str(int(v)))
                else:
                    cleaned.append(str(v).strip())

            is_header = 1 if r <= 5 and any("№" in str(x) or "Т/р" in str(x) or "номи" in str(x) or "йил" in str(x) for x in cleaned) else 0
            is_total = 1 if any("жами" in str(x).lower() for x in cleaned) else 0

            cursor.execute("""
            INSERT INTO excel_15_talik_sheets (sheet_index, sheet_name, row_index, cells_json, is_header, is_total)
            VALUES (?, ?, ?, ?, ?, ?)
            """, (s_idx, sname, r, json.dumps(cleaned, ensure_ascii=False), is_header, is_total))
            total_inserted += 1

    conn.commit()
    conn.close()
    wb.close()
    return True, f"Muvaffaqiyatli yuklandi: {len(wb.sheetnames)} ta sahifa, {total_inserted} ta qator yangilandi."

if __name__ == '__main__':
    init_15_talik_exact_cache()
